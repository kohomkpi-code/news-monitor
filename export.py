"""[프로] 탭의 조회 결과 내려받기 (엑셀 .xlsx 한 파일).

목적: 앱이 고른 결과(등급·점수·분류·키워드)와 담당자의 최종 선택을 함께 남겨,
Claude 등에 붙여 "점수 규칙·키워드·제외어를 어떻게 고치면 좋을까"를 물어볼 수 있게 한다.

시트 4개:
1) 조회결과   — 표에 나온 기사 전부(표 순서), 최종 선택 여부, 점수 항목별 내역
2) 빠진 기사  — 제외어·단순 행사·이미 보고·같은 사안 묶음으로 표에서 빠진 기사와 그 이유
3) 조회 조건  — 기간, 검색 방식, 옵션, 키워드·제외어 목록, 건수 흐름
4) 점수 규칙  — 이 버전의 점수 규칙과 리스크·정책 단어
네이버 Client ID/Secret 같은 비밀값은 넣지 않는다.
"""

import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import ranking
from naver_api import Article

HEAD_FILL = PatternFill("solid", fgColor="1F3B5C")
HEAD_FONT = Font(bold=True, color="FFFFFF")
A_FILL = PatternFill("solid", fgColor="FBF3E4")
MANUAL_FILL = PatternFill("solid", fgColor="E4F0F2")
PART_KEYS = ["등급", "키워드 순위", "여러 매체", "여러 키워드", "리스크 단어", "정책 단어", "전문지", "최근 24시간"]


def _sheet(ws, heads: list[tuple[str, int]], rows: list[list], fills: list | None = None):
    for c, (head, width) in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=head)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if fills and fills[r - 2]:
                cell.fill = fills[r - 2]
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(heads))}{len(rows) + 1}"


def _when(a: Article) -> str:
    return a.published.strftime("%Y-%m-%d %H:%M")


def save_xlsx(path: str, shown: list[tuple[Article, bool, bool, dict]], dropped: list[tuple[str, str, Article]],
              conditions: list[tuple[str, str]]) -> None:
    """shown = [(기사, 프로 선택, 베이직 선택, 점수 항목별 내역)] 표 순서대로. dropped = [(이유, 걸린 단어, 기사)]."""
    wb = Workbook()

    ws = wb.active
    ws.title = "조회결과"
    heads = [("순번", 6), ("프로 선택", 9), ("베이직 선택", 10), ("★추천", 7), ("등급", 6), ("점수", 6),
             ("AI 분류", 10), ("AI 인사이트", 30), ("키워드", 22), ("날짜", 16), ("매체", 16),
             ("묶인 기사 수", 9), ("제목", 60), ("요약", 50), ("링크", 40)]
    heads += [(f"점수:{k}", 9) for k in PART_KEYS] + [("걸린 리스크 단어", 14), ("걸린 정책 단어", 14)]
    rows, fills = [], []
    for n, (a, pro_on, basic_on, parts) in enumerate(shown, 1):
        rows.append([n, "O" if pro_on else "", "O" if basic_on else "", "★" if a.best else "", a.grade, a.score,
                     a.category, a.insight, ", ".join(a.keywords), _when(a), a.source, len(a.others),
                     a.title, a.description, a.link or a.naver_link]
                    + [parts.get(k, 0) for k in PART_KEYS]
                    + [", ".join(ranking.hit_words(a.title, ranking.RISK_WORDS)),
                       ", ".join(ranking.hit_words(a.title, ranking.POLICY_WORDS))])
        fills.append(MANUAL_FILL if a.grade == "직접" else A_FILL if a.grade == "A" else None)
    _sheet(ws, heads, rows, fills)

    ws = wb.create_sheet("빠진 기사")
    rows = [[reason, words, a.grade, ", ".join(a.keywords), _when(a), a.source, a.title, a.link or a.naver_link, ""]
            for reason, words, a in dropped]
    # 같은 사안으로 묶여 대표 기사 밑에 들어간 기사들 — 매체·제목·링크만 안다
    for n, (a, *_rest) in enumerate(shown, 1):
        for source, title, link in a.others:
            rows.append(["같은 사안 묶음", "", "", "", "", source, title, link, f"{n}번 기사에 묶임"])
    _sheet(ws, [("빠진 이유", 14), ("걸린 단어", 16), ("등급", 6), ("키워드", 22), ("날짜", 16), ("매체", 16),
                ("제목", 60), ("링크", 40), ("비고", 18)], rows)

    ws = wb.create_sheet("조회 조건")
    _sheet(ws, [("항목", 22), ("값", 100)], [list(c) for c in conditions])
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("점수 규칙")
    rules = [line.strip() for line in (ranking.__doc__ or "").splitlines() if line.strip()]
    rules += ["", "리스크 단어: " + ", ".join(ranking.RISK_WORDS), "정책 단어: " + ", ".join(ranking.POLICY_WORDS),
              "전문지 가점 매체: " + ", ".join(ranking.SPECIALIST_SOURCES),
              f"추천 고르기: 같은 묶음(공단 A 기사 / 대표 키워드)에서 {ranking.PER_KEYWORD}건까지"]
    _sheet(ws, [("규칙", 120)], [[r] for r in rules])

    wb.save(path)


def default_name(when: dt.datetime) -> str:
    return f"KOHOM뉴스룸_조회결과_{when:%Y%m%d_%H%M}.xlsx"
